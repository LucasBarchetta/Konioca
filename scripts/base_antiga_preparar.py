#!/usr/bin/env python3
"""Prepara a base antiga (exportação consolidada do Sults) para a tabela public.base_antiga.

Entrada: .xlsx (aba "Base consolidada") ou .csv (separador ; ou ,) com as colunas:
  Prioridade, Canal inicial, Nome, Primeiro nome, E-mail, WhatsApp (+55), Cidade, UF, Região,
  Primeiro contato, Último contato, Nº de cadastros, Gancho (mês do 1º contato), IDs no Sults, Observações

Regras (aprovadas):
  - P1 = WhatsApp + e-mail; P2 a P4 começam por e-mail. A coluna "Canal inicial" manda quando existe.
  - Sem telefone ou telefone incompleto (10 dígitos) não recebe WhatsApp.
  - "possível cadastro de teste" fica marcado para revisão humana e fora dos envios.

Uso:
  python3 scripts/base_antiga_preparar.py ARQUIVO                 # só o resumo, sem dados pessoais na tela
  python3 scripts/base_antiga_preparar.py ARQUIVO --saida x.json  # grava as linhas normalizadas (fora do repositório)
  python3 scripts/base_antiga_preparar.py ARQUIVO --enviar        # grava direto no Supabase (SUPABASE_URL, SUPABASE_SERVICE_ROLE_KEY)

Dados pessoais: nunca salvar a planilha nem o JSON dentro do repositório (.gitignore bloqueia *.xlsx e base_antiga*.json).
"""
import argparse, csv, collections, datetime as dt, json, os, re, sys, urllib.request

DDDS = {11,12,13,14,15,16,17,18,19,21,22,24,27,28,31,32,33,34,35,37,38,41,42,43,44,45,46,47,48,49,
        51,53,54,55,61,62,63,64,65,66,67,68,69,71,73,74,75,77,79,81,82,83,84,85,86,87,88,89,
        91,92,93,94,95,96,97,98,99}

CANAIS = {
    "whatsapp + e-mail": "whatsapp_email",
    "e-mail primeiro; whatsapp em lotes": "email_depois_whatsapp_lotes",
    "e-mail; whatsapp só se clicar": "email_whatsapp_se_clicar",
}

COLS = {
    "prioridade": "Prioridade", "canal": "Canal inicial", "nome": "Nome", "primeiro_nome": "Primeiro nome",
    "email": "E-mail", "whatsapp": "WhatsApp (+55)", "cidade": "Cidade", "uf": "UF", "regiao": "Região",
    "primeiro_contato": "Primeiro contato", "ultimo_contato": "Último contato", "n_cadastros": "Nº de cadastros",
    "gancho": "Gancho (mês do 1º contato)", "sults_ids": "IDs no Sults", "observacoes": "Observações",
}


def celular_e164(bruto):
    """Só celular completo: +55 + DDD válido + 9 + 8 dígitos. Qualquer outra coisa: None (sem WhatsApp)."""
    d = re.sub(r"\D", "", str(bruto or ""))
    if d.startswith("55") and len(d) == 13:
        d = d[2:]
    if len(d) != 11 or int(d[:2]) not in DDDS or d[2] != "9" or re.fullmatch(r"(\d)\1{8}", d[2:]):
        return None
    return "+55" + d


def texto(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def data_iso(v):
    if v in (None, ""):
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        d = v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)
    else:
        s = str(v).strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
            try:
                d = dt.datetime.strptime(s, fmt)
                break
            except ValueError:
                d = None
        if d is None:
            return None
    return d.strftime("%Y-%m-%dT%H:%M:%S-03:00")


def normalizar(linha, carga):
    g = lambda k: linha.get(COLS[k])
    obs = texto(g("observacoes")) or ""
    prio_rot = texto(g("prioridade")) or ""
    m = re.match(r"\s*(P[1-4])\b", prio_rot)
    prioridade = m.group(1) if m else None
    canal_rot = texto(g("canal")) or ""
    canal = CANAIS.get(canal_rot.lower()) or ("whatsapp_email" if prioridade == "P1" else "email_whatsapp_se_clicar")
    e164 = celular_e164(g("whatsapp"))
    sem_whats = "sem telefone" in obs.lower() or "dígitos" in obs.lower() or "digitos" in obs.lower()
    whatsapp_ok = bool(e164) and not sem_whats
    revisar = "cadastro de teste" in obs.lower()
    email = (texto(g("email")) or "").lower()
    n = g("n_cadastros")
    return {
        "email": email,
        "nome": texto(g("nome")) or texto(g("primeiro_nome")) or "",
        "primeiro_nome": texto(g("primeiro_nome")),
        "whatsapp_bruto": texto(g("whatsapp")),
        "whatsapp_e164": e164 if whatsapp_ok else None,
        "whatsapp_ok": whatsapp_ok,
        "cidade": texto(g("cidade")),
        "uf": texto(g("uf")),
        "regiao": texto(g("regiao")),
        "primeiro_contato": data_iso(g("primeiro_contato")),
        "ultimo_contato": data_iso(g("ultimo_contato")),
        "n_cadastros": int(n) if str(n or "").strip().isdigit() else None,
        "sults_ids": texto(g("sults_ids")),
        "prioridade": prioridade,
        "prioridade_rotulo": prio_rot or None,
        "canal_inicial": canal,
        "canal_inicial_rotulo": canal_rot or None,
        "gancho": texto(g("gancho")),
        "observacoes": obs or None,
        "revisar": revisar,
        "status": "revisar" if revisar else "importado",
        "carga": carga,
    }


def ler(caminho):
    if caminho.lower().endswith(".xlsx"):
        import openpyxl  # só para .xlsx
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        ws = wb["Base consolidada"] if "Base consolidada" in wb.sheetnames else wb.worksheets[0]
        linhas = list(ws.iter_rows(values_only=True))
        cab = [str(c).strip() if c is not None else "" for c in linhas[0]]
        return [dict(zip(cab, r)) for r in linhas[1:] if any(c not in (None, "") for c in r)]
    with open(caminho, encoding="utf-8-sig", newline="") as f:
        cab = f.readline(); f.seek(0)
        sep = ";" if cab.count(";") > cab.count(",") else ","
        return [dict(r) for r in csv.DictReader(f, delimiter=sep)]


def preparar(caminho):
    carga = f"{os.path.basename(caminho)}@{dt.date.today().isoformat()}"
    brutas = ler(caminho)
    faltando = [c for c in COLS.values() if brutas and c not in brutas[0]]
    if faltando:
        sys.exit(f"colunas ausentes: {faltando}")
    linhas, invalidas, vistos = [], 0, set()
    for b in brutas:
        n = normalizar(b, carga)
        if not n["prioridade"] or not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]{2,}", n["email"]) or n["email"] in vistos:
            invalidas += 1
            continue
        vistos.add(n["email"])
        linhas.append(n)
    return linhas, invalidas


def resumo(linhas, invalidas):
    c = collections.Counter((l["prioridade"], l["canal_inicial"]) for l in linhas)
    return {
        "linhas": len(linhas), "descartadas": invalidas,
        "por_prioridade_e_canal": {f"{p} · {k}": v for (p, k), v in sorted(c.items())},
        "com_whatsapp": sum(l["whatsapp_ok"] for l in linhas),
        "sem_whatsapp": sum(not l["whatsapp_ok"] for l in linhas),
        "para_revisar": sum(l["revisar"] for l in linhas),
    }


def enviar(linhas):
    url = os.environ["SUPABASE_URL"].rstrip("/") + "/rest/v1/base_antiga?on_conflict=email_norm"
    chave = os.environ["SUPABASE_SERVICE_ROLE_KEY"]
    for i in range(0, len(linhas), 500):
        req = urllib.request.Request(url, data=json.dumps(linhas[i:i + 500]).encode(), method="POST", headers={
            "apikey": chave, "authorization": f"Bearer {chave}", "content-type": "application/json",
            "prefer": "resolution=merge-duplicates,return=minimal"})
        urllib.request.urlopen(req).read()


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("arquivo")
    ap.add_argument("--saida")
    ap.add_argument("--enviar", action="store_true")
    a = ap.parse_args()
    linhas, invalidas = preparar(a.arquivo)
    print(json.dumps(resumo(linhas, invalidas), ensure_ascii=False, indent=2))
    if a.saida:
        with open(a.saida, "w", encoding="utf-8") as f:
            json.dump(linhas, f, ensure_ascii=False)
    if a.enviar:
        enviar(linhas)
        print("enviado:", len(linhas))
